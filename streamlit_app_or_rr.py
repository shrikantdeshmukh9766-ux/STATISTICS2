"""
Stream-lite — Odds Ratio / Relative Risk Builder
==================================================
A Streamlit app that turns an uploaded master chart (Excel/CSV) into a
publication-ready univariate OR / RR table, split by outcome group:
continuous factors get descriptive statistics per outcome group, categorical
factors get n (%) per outcome group, and every factor gets a univariate
odds ratio and/or relative risk (95% CI) with a p-value.

Run with:
    pip install streamlit pandas numpy scipy statsmodels openpyxl python-docx
    streamlit run streamlit_app_or_rr.py

How the analysis works
-----------------------
Outcome    -> pick a column, then pick which category is the "baseline"
              (reference / non-event) and which is the "event". If the
              column has more than two categories, rows with any other
              value are excluded from the analysis. The table has one
              column per outcome group (baseline, event).

Continuous (numeric) factors -> reported per outcome group as mean \u00B1 SD,
              median (IQR), or both — your choice, with a decimal-places
              control. "Auto" picks mean \u00B1 SD when both outcome groups
              pass a D'Agostino-Pearson normality test, median (IQR)
              otherwise.
                OR = exp(beta) from univariate logistic regression
                RR = exp(beta) from log-binomial regression, falling back
                     to modified Poisson regression with robust (HC1)
                     standard errors if log-binomial fails to converge.
              Effect size is reported per 1 unit or per 1 SD increase,
              your choice.

Categorical factors -> reported as n (%) per outcome group (% of that
              group's non-missing total for the variable). Pick a baseline
              (reference) category; every other category is compared
              pairwise against it using a 2x2 table:
                OR = (a*d) / (b*c)              [Woolf logit 95% CI]
                RR = risk(exposed) / risk(ref)   [log-method 95% CI]
                p  = chi-square test of independence, automatically
                     switched to Fisher's exact test when an expected
                     cell count is below 5.
              A 2x2 table with a zero cell gets the Haldane-Anscombe
              correction (+0.5 to all four cells) so OR/RR/CI can still be
              computed; this is flagged in the footnotes.

OR/RR values and descriptive statistics each have their own independent
decimal-places control in the settings panel.
"""

import io
import warnings
import numpy as np
import pandas as pd
import streamlit as st
from scipy import stats
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.styles import Font as OpenpyxlFont

try:
    import statsmodels.api as sm
except ImportError:
    sm = None

st.set_page_config(page_title="Stream-lite · OR / RR Builder", layout="wide")
warnings.filterwarnings("ignore")  # keep statsmodels convergence/domain chatter out of the app

# --------------------------------------------------------------------------
# Shared helpers
# --------------------------------------------------------------------------

def detect_type(series: pd.Series):
    """Guess whether a column is numerical or categorical."""
    non_missing = series.dropna()
    non_missing = non_missing[non_missing.astype(str).str.strip() != ""]
    n = len(non_missing)
    if n == 0:
        return "categorical", 0, 0
    numeric_coerced = pd.to_numeric(non_missing, errors="coerce")
    numeric_ratio = numeric_coerced.notna().mean()
    unique_n = non_missing.astype(str).str.strip().nunique()
    if numeric_ratio >= 0.9 and unique_n > 10:
        return "numerical", n, unique_n
    return "categorical", n, unique_n


def effective_type(meta):
    return meta["detected"] if meta["type"] == "auto" else meta["type"]


def is_normal(arr, alpha=0.05):
    """D'Agostino-Pearson omnibus normality test. Needs n>=8; smaller
    samples are treated as non-normal (safer default)."""
    arr = np.asarray(arr, dtype=float)
    arr = arr[~np.isnan(arr)]
    if len(arr) < 8:
        return False
    if np.all(arr == arr[0]):
        return True
    try:
        _, p = stats.normaltest(arr)
        return p > alpha
    except Exception:
        return False


def fmt_p(p):
    if p is None or (isinstance(p, float) and np.isnan(p)):
        return "—"
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def fmt_num(x, d=3):
    return f"{x:.{d}f}"


def fmt_ratio(val, low, high, decimals=2):
    return f"{val:.{decimals}f} ({low:.{decimals}f}\u2013{high:.{decimals}f})"


def format_numeric_cell(values, display_mode, use_param, decimals):
    """Format one outcome group's numeric summary. display_mode:
    'auto' | 'mean_sd' | 'median_iqr' | 'both'"""
    if len(values) == 0:
        return "—"
    show_mean = display_mode == "mean_sd" or display_mode == "both" or (display_mode == "auto" and use_param)
    show_median = display_mode == "median_iqr" or display_mode == "both" or (display_mode == "auto" and not use_param)
    parts = []
    if show_mean:
        parts.append(f"{fmt_num(np.mean(values), decimals)} \u00B1 {fmt_num(np.std(values, ddof=1), decimals)}")
    if show_median:
        q1, med, q3 = np.percentile(values, [25, 50, 75])
        parts.append(f"{fmt_num(med, decimals)} ({fmt_num(q1, decimals)}\u2013{fmt_num(q3, decimals)})")
    return "; ".join(parts)


def numeric_label(col, display_mode, use_param):
    if display_mode == "mean_sd":
        return f"{col}, mean \u00B1 SD"
    if display_mode == "median_iqr":
        return f"{col}, median (IQR)"
    if display_mode == "both":
        return f"{col}, mean \u00B1 SD; median (IQR)"
    return f"{col}, mean \u00B1 SD" if use_param else f"{col}, median (IQR)"


# --------------------------------------------------------------------------
# 2x2-table statistics (categorical factors)
# --------------------------------------------------------------------------

def chi_or_fisher_p(table, yates_correction=False):
    """RxC contingency table -> chi-square, auto-falling back to Fisher's
    exact test for 2x2 tables with low expected counts."""
    arr = np.array(table, dtype=float)
    chi2, p, dof, expected = stats.chi2_contingency(arr, correction=yates_correction)
    min_e = float(expected.min())
    if arr.shape == (2, 2) and min_e < 5:
        _, p_fisher = stats.fisher_exact(arr)
        return {"name": "Fisher's exact test", "p": float(p_fisher), "min_expected": min_e}
    name = "Chi-square test"
    if arr.shape == (2, 2) and yates_correction:
        name = "Chi-square test (Yates-corrected)"
    return {"name": name, "p": float(p), "min_expected": min_e}


def _haldane(a, b, c, d):
    if min(a, b, c, d) == 0:
        return a + 0.5, b + 0.5, c + 0.5, d + 0.5, True
    return a, b, c, d, False


def odds_ratio_ci(a, b, c, d, alpha):
    a, b, c, d, corrected = _haldane(a, b, c, d)
    or_val = (a * d) / (b * c)
    se = np.sqrt(1 / a + 1 / b + 1 / c + 1 / d)
    z = stats.norm.ppf(1 - alpha / 2)
    low = np.exp(np.log(or_val) - z * se)
    high = np.exp(np.log(or_val) + z * se)
    return {"val": or_val, "low": low, "high": high, "corrected": corrected}


def relative_risk_ci(a, b, c, d, alpha):
    a, b, c, d, corrected = _haldane(a, b, c, d)
    risk_exp = a / (a + b)
    risk_ref = c / (c + d)
    rr = risk_exp / risk_ref
    se = np.sqrt(1 / a - 1 / (a + b) + 1 / c - 1 / (c + d))
    z = stats.norm.ppf(1 - alpha / 2)
    low = np.exp(np.log(rr) - z * se)
    high = np.exp(np.log(rr) + z * se)
    return {"val": rr, "low": low, "high": high, "corrected": corrected}


# --------------------------------------------------------------------------
# Continuous-predictor statistics (numeric factors)
# --------------------------------------------------------------------------

def logistic_or(x, y, alpha, standardize=False):
    if sm is None:
        return {"error": "statsmodels is not installed"}
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    if standardize:
        sd = x.std(ddof=1)
        if sd > 0:
            x = (x - x.mean()) / sd
    X = sm.add_constant(x)
    try:
        model = sm.Logit(y, X).fit(disp=0)
        coef, se, p = model.params[1], model.bse[1], model.pvalues[1]
        z = stats.norm.ppf(1 - alpha / 2)
        return {"val": np.exp(coef), "low": np.exp(coef - z * se),
                "high": np.exp(coef + z * se), "p": float(p), "error": None}
    except Exception as e:
        return {"error": str(e)}


def poisson_rr(x, y, alpha, standardize=False):
    if sm is None:
        return {"error": "statsmodels is not installed"}
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    if standardize:
        sd = x.std(ddof=1)
        if sd > 0:
            x = (x - x.mean()) / sd
    X = sm.add_constant(x)
    z = stats.norm.ppf(1 - alpha / 2)
    method = "log-binomial"
    try:
        model = sm.GLM(y, X, family=sm.families.Binomial(link=sm.families.links.Log())).fit()
        coef, se = model.params[1], model.bse[1]
    except Exception:
        try:
            model = sm.GLM(y, X, family=sm.families.Poisson()).fit(cov_type="HC1")
            coef, se = model.params[1], model.bse[1]
            method = "modified Poisson (robust SE)"
        except Exception as e:
            return {"error": str(e)}
    p = float(2 * (1 - stats.norm.cdf(abs(coef / se)))) if se > 0 else float("nan")
    return {"val": np.exp(coef), "low": np.exp(coef - z * se),
            "high": np.exp(coef + z * se), "p": p, "method": method, "error": None}


# --------------------------------------------------------------------------
# Table builder
# --------------------------------------------------------------------------

def build_split_table(df, outcome_col, baseline_outcome, event_outcome, factor_cols,
                       factor_types, ref_map, alpha, compute_or, compute_rr, numeric_effect,
                       display_mode, desc_decimals, or_decimals, pct_digits, yates_correction,
                       pct_mode="column"):
    outcome_raw = df[outcome_col].astype(str).str.strip()
    in_scope = outcome_raw.isin([baseline_outcome, event_outcome])
    y_full = pd.Series(np.nan, index=df.index)
    y_full[in_scope & (outcome_raw == event_outcome)] = 1.0
    y_full[in_scope & (outcome_raw == baseline_outcome)] = 0.0

    n_baseline_all = int((y_full == 0).sum())
    n_event_all = int((y_full == 1).sum())

    ci_pct = int(round((1 - alpha) * 100))
    header = ["Variable", f"{event_outcome} (n={n_event_all})", f"{baseline_outcome} (n={n_baseline_all})"]
    if compute_or:
        header.append(f"OR ({ci_pct}% CI)")
    if compute_rr:
        header.append(f"RR ({ci_pct}% CI)")
    header.append("p-value")

    csv_rows = [header]
    display_rows = []
    flags = set()

    for col in factor_cols:
        vtype = factor_types[col]

        if vtype == "categorical":
            series = df[col].astype(str).str.strip()
            series = series.where(df[col].notna() & (series != ""), other=np.nan)
            levels = sorted(series.dropna().unique().tolist())
            ref = ref_map.get(col)
            if ref not in levels:
                continue

            base_total_var = int(((series.notna()) & (y_full == 0)).sum())
            event_total_var = int(((series.notna()) & (y_full == 1)).sum())

            display_rows.append({"kind": "varheader", "label": f"{col}, n (%) (ref: {ref})"})
            csv_rows.append([f"{col}, n (%) (ref: {ref})"])

            ref_mask = (series == ref)
            c = int((ref_mask & (y_full == 1)).sum())
            d = int((ref_mask & (y_full == 0)).sum())
            if pct_mode == "row":
                row_total_ref = c + d
                pct_d = 100 * d / row_total_ref if row_total_ref else 0.0
                pct_c = 100 * c / row_total_ref if row_total_ref else 0.0
            else:
                pct_d = 100 * d / base_total_var if base_total_var else 0.0
                pct_c = 100 * c / event_total_var if event_total_var else 0.0
            ref_cells = [f"{c} ({pct_c:.{pct_digits}f}%)", f"{d} ({pct_d:.{pct_digits}f}%)"]
            if compute_or:
                ref_cells.append("1.00 (Reference)")
            if compute_rr:
                ref_cells.append("1.00 (Reference)")
            ref_cells.append("—")
            display_rows.append({"kind": "level", "label": ref, "cells": ref_cells, "sig": False})
            csv_rows.append([f"  {ref}", *ref_cells])

            for lv in levels:
                if lv == ref:
                    continue
                lv_mask = (series == lv)
                a = int((lv_mask & (y_full == 1)).sum())
                b = int((lv_mask & (y_full == 0)).sum())
                if pct_mode == "row":
                    row_total_lv = a + b
                    pct_b = 100 * b / row_total_lv if row_total_lv else 0.0
                    pct_a = 100 * a / row_total_lv if row_total_lv else 0.0
                else:
                    pct_b = 100 * b / base_total_var if base_total_var else 0.0
                    pct_a = 100 * a / event_total_var if event_total_var else 0.0
                cells = [f"{a} ({pct_a:.{pct_digits}f}%)", f"{b} ({pct_b:.{pct_digits}f}%)"]
                p_val = None
                if (a + b) == 0 or (c + d) == 0:
                    if compute_or:
                        cells.append("—")
                    if compute_rr:
                        cells.append("—")
                    cells.append("—")
                    flags.add("skipped")
                else:
                    test_res = chi_or_fisher_p([[a, b], [c, d]], yates_correction=yates_correction)
                    p_val = test_res["p"]
                    if test_res["name"].startswith("Fisher"):
                        flags.add("fisher")
                    else:
                        flags.add("chi2")
                        if test_res["min_expected"] < 5:
                            flags.add("lowE")
                    if compute_or:
                        r = odds_ratio_ci(a, b, c, d, alpha)
                        if r["corrected"]:
                            flags.add("haldane")
                        cells.append(fmt_ratio(r["val"], r["low"], r["high"], or_decimals))
                    if compute_rr:
                        r = relative_risk_ci(a, b, c, d, alpha)
                        if r["corrected"]:
                            flags.add("haldane")
                        cells.append(fmt_ratio(r["val"], r["low"], r["high"], or_decimals))
                    cells.append(fmt_p(p_val))
                sig = p_val is not None and p_val < alpha
                display_rows.append({"kind": "level", "label": lv, "cells": cells, "sig": sig})
                csv_rows.append([f"  {lv}", *cells])

        else:  # numerical
            numeric_series = pd.to_numeric(df[col], errors="coerce")
            baseline_vals = numeric_series[(y_full == 0) & numeric_series.notna()].values
            event_vals = numeric_series[(y_full == 1) & numeric_series.notna()].values

            if display_mode == "auto":
                use_param = is_normal(baseline_vals, alpha) and is_normal(event_vals, alpha)
            else:
                use_param = display_mode == "mean_sd"

            label = numeric_label(col, display_mode, use_param)
            cell_baseline = format_numeric_cell(baseline_vals, display_mode, use_param, desc_decimals)
            cell_event = format_numeric_cell(event_vals, display_mode, use_param, desc_decimals)
            cells = [cell_event, cell_baseline]

            mask = numeric_series.notna() & y_full.notna()
            xv = numeric_series[mask].values
            yv = y_full[mask].values
            n = len(xv)

            p_val = None
            if n < 8 or len(np.unique(yv)) < 2:
                if compute_or:
                    cells.append("—")
                if compute_rr:
                    cells.append("—")
                cells.append("—")
                flags.add("skipped")
            else:
                if compute_or:
                    lg = logistic_or(xv, yv, alpha, standardize=(numeric_effect == "sd"))
                    if lg.get("error"):
                        cells.append("—")
                        flags.add("skipped")
                    else:
                        cells.append(fmt_ratio(lg["val"], lg["low"], lg["high"], or_decimals))
                        p_val = lg["p"]
                if compute_rr:
                    ps = poisson_rr(xv, yv, alpha, standardize=(numeric_effect == "sd"))
                    if ps.get("error"):
                        cells.append("—")
                        flags.add("skipped")
                    else:
                        cells.append(fmt_ratio(ps["val"], ps["low"], ps["high"], or_decimals))
                        if p_val is None:
                            p_val = ps["p"]
                        if ps.get("method") == "modified Poisson (robust SE)":
                            flags.add("modpoisson")
                cells.append(fmt_p(p_val) if p_val is not None else "—")
            sig = p_val is not None and p_val < alpha
            display_rows.append({"kind": "var", "label": label, "cells": cells, "sig": sig})
            csv_rows.append([label, *cells])

    return header, display_rows, csv_rows, flags


def build_footnotes(alpha, flags, outcome_col, baseline_outcome, event_outcome, n_excluded,
                     yates_correction, compute_or, compute_rr, display_mode, desc_decimals,
                     or_decimals, numeric_effect, pct_digits, pct_mode="column"):
    ci_pct = int(round((1 - alpha) * 100))
    notes = []

    desc_txt = {"mean_sd": "mean \u00B1 SD", "median_iqr": "median (IQR)",
                "both": "mean \u00B1 SD and median (IQR)",
                "auto": f"mean \u00B1 SD (assessed as normal via D'Agostino-Pearson test, \u03B1={alpha}) "
                        "or median (IQR) otherwise"}[display_mode]
    if pct_mode == "row":
        pct_basis = ("row-wise: each n (%) is a share of that category's total across both outcome "
                     "groups (rows sum to ~100%)")
    else:
        pct_basis = ("column-wise: each n (%) is a share of its own outcome group's non-missing total "
                     "for that variable (columns sum to ~100%)")
    notes.append(f"Continuous variables reported as {desc_txt}, to {desc_decimals} decimal place(s), per "
                  f"outcome group. Categorical variables reported as n (%), {pct_basis}, to {pct_digits} "
                  f"decimal place(s) (missing values excluded from the denominator).")

    parts = []
    if compute_or:
        parts.append("OR = odds ratio")
    if compute_rr:
        parts.append("RR = relative risk (risk ratio)")
    notes.append("; ".join(parts) + f"; CI = confidence interval ({ci_pct}%), shown to {or_decimals} "
                  f"decimal place(s).")

    notes.append(
        f"Outcome variable: {outcome_col}. Baseline (reference) group: {baseline_outcome}. "
        f"Event group: {event_outcome}."
        + (f" {n_excluded} row(s) with another {outcome_col} value were excluded from the analysis."
           if n_excluded > 0 else "")
    )

    cat_note = ("Categorical factors: OR/RR computed pairwise against the stated baseline category, "
                "restricted to rows in the compared category or the reference category; p-value from "
                "chi-square test of independence")
    cat_note += (" (Yates' continuity correction applied to 2\u00D72 tables)" if yates_correction
                 else " (no continuity correction applied)")
    cat_note += ", automatically switched to Fisher's exact test when an expected cell count is below 5."
    notes.append(cat_note)

    if "haldane" in flags:
        notes.append("Haldane-Anscombe correction (0.5 added to all four cell counts) applied where a "
                      "2\u00D72 table contained a zero cell, to allow OR/RR and CI calculation.")
    if "lowE" in flags:
        notes.append("Caution: one or more chi-square comparisons above have expected cell counts below "
                      "5; the chi-square approximation may be unreliable there.")

    num_parts = []
    eff_txt = "per 1 SD increase" if numeric_effect == "sd" else "per 1 unit increase"
    if compute_or:
        num_parts.append(f"OR ({eff_txt}) from univariate logistic regression (Wald {ci_pct}% CI)")
    if compute_rr:
        rr_txt = f"RR ({eff_txt}) from log-binomial regression"
        if "modpoisson" in flags:
            rr_txt += ", falling back to modified Poisson regression with robust (HC1) standard errors when log-binomial failed to converge"
        num_parts.append(rr_txt)
    if num_parts:
        notes.append("Numeric factors: " + "; ".join(num_parts) + ".")

    if "skipped" in flags:
        notes.append("A comparison could not be computed for one or more categories/variables (e.g. "
                      "insufficient data, zero-variance predictor, or model non-convergence) and was "
                      "left blank.")
    notes.append(f"Bold p-values indicate statistical significance at \u03B1={alpha}.")
    return notes


# --------------------------------------------------------------------------
# Rendering / export
# --------------------------------------------------------------------------

def render_table_markdown(header, display_rows):
    css = """
    <style>
    table.pub { width:100%; border-collapse:collapse; font-family: Calibri, Candara, Segoe, "Segoe UI", Optima, Arial, sans-serif; font-size: 9pt; }
    table.pub thead th { border-top:2px solid #1E2A32; border-bottom:1px solid #1E2A32;
                          padding:8px 10px; text-align:left; font-family: inherit; }
    table.pub tbody td { padding:5px 10px; }
    table.pub tbody tr.var td { font-weight:700; padding-top:10px; }
    table.pub tbody tr.level td.name { padding-left:20px; color:#555; font-weight:400; }
    table.pub td.stat { font-family: inherit; text-align:center; font-size:9pt;}
    table.pub td.sig { font-weight:700; }
    table.pub tbody tr.lastrow td { border-bottom:2px solid #1E2A32; padding-bottom:10px; }
    </style>
    """
    html = css + '<table class="pub"><thead><tr>'
    for h in header:
        html += f"<th>{h}</th>"
    html += "</tr></thead><tbody>"

    for idx, row in enumerate(display_rows):
        is_last_of_block = (idx == len(display_rows) - 1) or \
                            (display_rows[idx + 1]["kind"] in ("var", "varheader"))
        cls = "var" if row["kind"] in ("var", "varheader") else "level"
        cls += " lastrow" if is_last_of_block else ""
        html += f'<tr class="{cls}">'

        if row["kind"] == "varheader":
            html += f'<td>{row["label"]}</td>'
            html += f'<td colspan="{len(header) - 1}"></td>'
        else:
            name_cls = "name" if row["kind"] == "level" else ""
            html += f'<td class="{name_cls}">{row["label"]}</td>'
            n_cells = len(row["cells"])
            for i, c in enumerate(row["cells"]):
                is_p = i == n_cells - 1
                cls2 = "stat sig" if (is_p and row.get("sig")) else "stat"
                html += f'<td class="{cls2}">{c}</td>'
        html += "</tr>"
    html += "</tbody></table>"
    return html


def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge in kwargs:
            spec = kwargs[edge]
            tag = f'w:{edge}'
            el = tcBorders.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                tcBorders.append(el)
            el.set(qn('w:val'), spec.get('val', 'single'))
            el.set(qn('w:sz'), str(spec.get('sz', 8)))
            el.set(qn('w:color'), spec.get('color', '000000'))


def _set_run_font(run, name="Calibri", size=9, bold=None, italic=None):
    run.font.name = name
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def build_excel(csv_rows, sheet_name="OR-RR Table"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.sheet_state = "visible"
    for row in csv_rows:
        ws.append(row)
    for xl_row in ws.iter_rows():
        for cell in xl_row:
            bold = cell.row == 1
            cell.font = OpenpyxlFont(name="Calibri", size=9, bold=bold)
    for column_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 45)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def build_docx(header, display_rows, alpha, footnotes, title="Odds Ratio / Relative Risk Table"):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9)

    h = doc.add_heading(title, level=2)
    for r in h.runs:
        _set_run_font(r, size=11, bold=True)

    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.autofit = True

    hdr_cells = table.rows[0].cells
    for i, htext in enumerate(header):
        hdr_cells[i].text = str(htext)
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                _set_run_font(run, size=9, bold=True)
        _set_cell_border(hdr_cells[i], top={'sz': 12, 'val': 'single'}, bottom={'sz': 8, 'val': 'single'})

    n_rows = len(display_rows)
    for idx, row in enumerate(display_rows):
        cells = table.add_row().cells
        is_last = idx == n_rows - 1
        next_is_new_block = is_last or (display_rows[idx + 1]["kind"] in ("var", "varheader"))

        if row["kind"] == "varheader":
            cells[0].text = row["label"]
            for run in cells[0].paragraphs[0].runs:
                _set_run_font(run, size=9, bold=True)
            for c in cells[1:]:
                c.text = ""
        else:
            label = ("    " + row["label"]) if row["kind"] == "level" else row["label"]
            cells[0].text = label
            for run in cells[0].paragraphs[0].runs:
                _set_run_font(run, size=9, bold=(row["kind"] == "var"))
            n_cells = len(row["cells"])
            for ci, val in enumerate(row["cells"], start=1):
                cells[ci].text = str(val)
                is_p = ci == n_cells
                for p in cells[ci].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        _set_run_font(run, size=9, bold=(is_p and row.get("sig")))

        if next_is_new_block:
            for c in cells:
                _set_cell_border(c, bottom={'sz': 8, 'val': 'single'})

    for c in table.rows[-1].cells:
        _set_cell_border(c, bottom={'sz': 12, 'val': 'single'})

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if run.font.size is None:
                        _set_run_font(run, size=9)

    doc.add_paragraph()
    for f in footnotes:
        p = doc.add_paragraph(f)
        for run in p.runs:
            _set_run_font(run, size=9, italic=True)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Streamlit UI
# --------------------------------------------------------------------------

st.title("Stream-lite · OR / RR Builder")
st.caption(
    "Upload a master chart, pick an outcome and its baseline category, pick your factors and their "
    "baseline categories, and Stream-lite builds a split table: continuous descriptive stats or "
    "categorical n (%) per outcome group, plus univariate OR and/or RR (95% CI) and a p-value for "
    "every factor."
)

if sm is None:
    st.warning("`statsmodels` is not installed, so numeric factors can't be analyzed (categorical "
               "factors still work). Install it with `pip install statsmodels` and restart the app.")

st.markdown("### 1. Upload master chart")
uploaded = st.file_uploader("Excel (.xlsx/.xls) or CSV. First row must be column headers.",
                             type=["xlsx", "xls", "csv"])

if uploaded is not None:
    is_excel = not uploaded.name.lower().endswith(".csv")
    sheet_name = None
    try:
        if is_excel:
            xls = pd.ExcelFile(uploaded)
            sheet_names = xls.sheet_names
            sheet_name = st.selectbox("Select sheet", options=sheet_names) if len(sheet_names) > 1 else sheet_names[0]
            df = pd.read_excel(xls, sheet_name=sheet_name)
        else:
            df = pd.read_csv(uploaded)
    except Exception as e:
        st.error(f"Could not read that file: {e}")
        st.stop()

    st.success(f"Loaded **{uploaded.name}**"
                + (f" · sheet **{sheet_name}**" if sheet_name else "")
                + f" — {len(df)} rows, {len(df.columns)} columns")

    dataset_key = f"{uploaded.name}::{sheet_name}"
    if "var_meta" not in st.session_state or st.session_state.get("_last_file") != dataset_key:
        var_meta = {}
        total_rows = len(df)
        for col in df.columns:
            vtype, n, unique_n = detect_type(df[col])
            var_meta[col] = {"type": "auto", "detected": vtype, "use": False,
                              "n": n, "missing": total_rows - n, "unique": unique_n}
        st.session_state["var_meta"] = var_meta
        st.session_state["_last_file"] = dataset_key

    var_meta = st.session_state["var_meta"]

    # ---------------------------------------------------------------- 2. Outcome
    st.markdown("### 2. Outcome variable")
    outcome_col = st.selectbox("Outcome column", options=list(df.columns), key="outcome_col")

    out_series = df[outcome_col].astype(str).str.strip()
    out_series = out_series.where(df[outcome_col].notna() & (out_series != ""), other=np.nan)
    out_levels = sorted(out_series.dropna().unique().tolist())

    if len(out_levels) < 2:
        st.error(f"**{outcome_col}** has fewer than 2 distinct non-missing values — pick another column.")
        st.stop()

    counts = out_series.value_counts()
    default_baseline = counts.idxmax() if len(counts) else out_levels[0]
    ocol1, ocol2 = st.columns(2)
    with ocol1:
        baseline_outcome = st.selectbox(
            "Baseline (reference / non-event) category",
            options=out_levels, index=out_levels.index(default_baseline),
        )
    with ocol2:
        event_options = [lv for lv in out_levels if lv != baseline_outcome]
        event_outcome = st.selectbox("Event (outcome-positive) category", options=event_options)

    n_baseline = int((out_series == baseline_outcome).sum())
    n_event = int((out_series == event_outcome).sum())
    n_excluded = int((out_series.notna() & ~out_series.isin([baseline_outcome, event_outcome])).sum())

    info = f"Baseline **{baseline_outcome}**: n={n_baseline} · Event **{event_outcome}**: n={n_event}"
    if n_excluded:
        info += f" · {n_excluded} row(s) with other {outcome_col} values excluded from analysis"
    st.info(info)
    if min(n_baseline, n_event) < 5:
        st.warning("One outcome category has fewer than 5 observations — OR/RR estimates will be unstable.")

    # ---------------------------------------------------------------- 3. Factors
    st.markdown("### 3. Factors")
    st.caption("Check **Use** for each variable you want analyzed as a factor. The outcome column is "
               "excluded automatically.")

    factor_candidates = [c for c in df.columns if c != outcome_col]
    editor_df = pd.DataFrame([
        {
            "Variable": col,
            "Use": var_meta[col]["use"],
            "Type": var_meta[col]["type"],
            "Auto-detected": var_meta[col]["detected"].capitalize(),
            "n (non-missing)": var_meta[col]["n"],
            "n (missing)": var_meta[col]["missing"],
            "Unique values": var_meta[col]["unique"],
        }
        for col in factor_candidates
    ])

    edited = st.data_editor(
        editor_df,
        column_config={
            "Use": st.column_config.CheckboxColumn(required=True),
            "Type": st.column_config.SelectboxColumn(options=["auto", "numerical", "categorical"], required=True),
            "Auto-detected": st.column_config.TextColumn(disabled=True),
            "Variable": st.column_config.TextColumn(disabled=True),
            "n (non-missing)": st.column_config.NumberColumn(disabled=True),
            "n (missing)": st.column_config.NumberColumn(disabled=True),
            "Unique values": st.column_config.NumberColumn(disabled=True),
        },
        hide_index=True,
        use_container_width=True,
        key="factor_editor",
    )
    for _, row in edited.iterrows():
        var_meta[row["Variable"]]["use"] = bool(row["Use"])
        var_meta[row["Variable"]]["type"] = row["Type"]

    selected_factors = [c for c in factor_candidates if var_meta[c]["use"]]
    factor_types = {c: effective_type(var_meta[c]) for c in selected_factors}
    categorical_factors = [c for c in selected_factors if factor_types[c] == "categorical"]
    numeric_factors = [c for c in selected_factors if factor_types[c] == "numerical"]

    ref_map = {}
    if categorical_factors:
        st.markdown("**Baseline (reference) category per categorical factor**")
        for col in categorical_factors:
            series = df[col].astype(str).str.strip()
            series = series.where(df[col].notna() & (series != ""), other=np.nan)
            levels = sorted(series.dropna().unique().tolist())
            if len(levels) < 2:
                st.warning(f"**{col}** has fewer than 2 categories — it will be skipped.")
                continue
            default_ref = series.value_counts().idxmax()
            rc1, rc2 = st.columns([2, 3])
            with rc1:
                st.markdown(f"`{col}`")
            with rc2:
                ref_map[col] = st.selectbox(
                    f"Baseline for {col}", options=levels,
                    index=levels.index(default_ref), key=f"ref_{col}",
                    label_visibility="collapsed",
                )

    # ---------------------------------------------------------------- 4. Settings
    st.markdown("### 4. Analysis settings")
    s1, s2, s3, s4 = st.columns([1.4, 1.4, 1.1, 1.1])
    with s1:
        st.markdown("**Effect measure**")
        compute_or = st.checkbox("Compute Odds Ratio (OR)", value=True)
        compute_rr = st.checkbox("Compute Relative Risk (RR)", value=True)
        or_decimals = st.number_input("OR/RR decimal places", min_value=0, max_value=6, value=2, step=1)
    with s2:
        st.markdown("**Continuous descriptive stats**")
        display_mode = st.radio(
            "Display as",
            options=["auto", "mean_sd", "median_iqr", "both"],
            format_func=lambda x: {"auto": "Auto (normality-based)", "mean_sd": "Mean \u00B1 SD",
                                    "median_iqr": "Median (IQR)", "both": "Both"}[x],
            label_visibility="collapsed",
        )
        desc_decimals = st.number_input("Descriptive stats decimal places", min_value=0, max_value=6, value=3, step=1)
    with s3:
        numeric_effect = st.radio(
            "Numeric factor effect size",
            options=["unit", "sd"],
            format_func=lambda x: "Per 1 unit increase" if x == "unit" else "Per 1 SD increase",
        )
        alpha = st.number_input("Significance level (\u03B1)", min_value=0.001, max_value=0.5, value=0.05, step=0.01)
    with s4:
        pct_mode = st.radio(
            "Categorical % basis",
            options=["column", "row"],
            format_func=lambda x: "Column-wise (\u00F7 group n)" if x == "column" else "Row-wise (\u00F7 category n)",
            help="Column-wise: each n (%) is a share of its own outcome group's total (columns sum to "
                 "~100%). Row-wise: each n (%) is a share of that category's total across both outcome "
                 "groups (rows sum to ~100%).",
        )
        pct_digits = st.number_input("Categorical % decimal places", min_value=0, max_value=4, value=1, step=1)
        yates_correction = st.checkbox(
            "Apply Yates' continuity correction (2\u00D72 chi-square)", value=False,
            help="Only affects chi-square p-values on 2\u00D72 tables (ignored for Fisher's exact test). "
                 "Off by default.",
        )

    can_generate = True
    if not selected_factors:
        st.warning("No factors selected — check **Use** for at least one variable above.")
        can_generate = False
    if not (compute_or or compute_rr):
        st.warning("Select at least one of Odds Ratio or Relative Risk.")
        can_generate = False
    if numeric_factors and sm is None:
        st.warning("Numeric factors are selected but `statsmodels` isn't installed — those rows will be skipped.")

    if st.button("Generate table", type="primary", disabled=not can_generate):
        header, display_rows, csv_rows, flags = build_split_table(
            df, outcome_col, baseline_outcome, event_outcome, selected_factors, factor_types,
            ref_map, alpha, compute_or, compute_rr, numeric_effect, display_mode, desc_decimals,
            or_decimals, pct_digits, yates_correction, pct_mode,
        )
        st.session_state["or_rr_result"] = (header, display_rows, csv_rows, flags, alpha,
                                             outcome_col, baseline_outcome, event_outcome, n_excluded,
                                             yates_correction, compute_or, compute_rr, display_mode,
                                             desc_decimals, or_decimals, numeric_effect, pct_digits,
                                             pct_mode)

    if "or_rr_result" in st.session_state:
        (header, display_rows, csv_rows, flags, r_alpha, r_outcome_col, r_baseline, r_event,
         r_n_excluded, r_yates, r_compute_or, r_compute_rr, r_display_mode, r_desc_decimals,
         r_or_decimals, r_numeric_effect, r_pct_digits, r_pct_mode) = st.session_state["or_rr_result"]

        st.markdown("### Table. Baseline characteristics with univariate OR / RR")
        st.markdown(render_table_markdown(header, display_rows), unsafe_allow_html=True)

        footnotes = build_footnotes(r_alpha, flags, r_outcome_col, r_baseline, r_event, r_n_excluded,
                                     r_yates, r_compute_or, r_compute_rr, r_display_mode, r_desc_decimals,
                                     r_or_decimals, r_numeric_effect, r_pct_digits, r_pct_mode)
        st.caption("  \n".join(footnotes))

        dl_col1, dl_col2, dl_col3 = st.columns(3)
        with dl_col1:
            csv_buf = io.StringIO()
            pd.DataFrame(csv_rows).to_csv(csv_buf, index=False, header=False)
            st.download_button("Download CSV", csv_buf.getvalue(), file_name="or_rr_table.csv", mime="text/csv")
        with dl_col2:
            excel_bytes = build_excel(csv_rows)
            st.download_button("Download Excel", excel_bytes, file_name="or_rr_table.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        with dl_col3:
            docx_bytes = build_docx(header, display_rows, r_alpha, footnotes)
            st.download_button("Download Word", docx_bytes, file_name="or_rr_table.docx",
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

else:
    st.info("Upload a file to get started. Nothing leaves your machine — the app runs locally.")
